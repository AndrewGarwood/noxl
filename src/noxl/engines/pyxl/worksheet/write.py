"""
*`@file`* `engines/pyxl/worksheet/write`
"""
from __future__ import annotations
from pathlib import Path
from typing import (Any, Optional, Mapping, Callable, Sequence, OrderedDict)
import lute
from pyqi import (Array, String, Function)
from openpyxl.utils import (get_column_letter)
from noxl.models import (CellAddress, CellAddressLike, HeaderIndex, normalize_cell_value)
from ..types import (PyxlWorkbook, PyxlWorksheet, PyxlWorkbookSource)
from ..utils import (resolve_pyxl_wb, get_visible_sheetnames)

LogMsg = lute.LogMsg
mlog = lute.get_logger(Path(__file__).parent.stem)

def delete_rows(
    ws: PyxlWorksheet, 
    /,
    row_min: int=2, 
    row_max: Optional[int]=None
) -> PyxlWorksheet:
    row_max = row_max if isinstance(row_max, int) else ws.max_row
    ws.delete_rows(row_min, row_max)
    return ws

def clear_range(
    ws: PyxlWorksheet,
    /,
    row_min: int = 2,
    row_max: Optional[int] = None,
    col_min: int = 1,
    col_max: Optional[int] = None
) -> PyxlWorksheet:
    """Sets the value of all cells within the specified range to None."""
    row_max = row_max if isinstance(row_max, int) else ws.max_row
    col_max = col_max if isinstance(col_max, int) else ws.max_column

    # openpyxl uses 1-based indexing; ensure bounds are valid
    for row in range(row_min, row_max + 1):
        for col in range(col_min, col_max + 1):
            ws.cell(row=row, column=col).value = None
    return ws

def write_rows(
    ws: PyxlWorksheet,
    rows: Sequence[Sequence[Any]],
    origin: CellAddressLike="A2",
    /,
) -> PyxlWorksheet:
    """ `O(M x N)` """
    try:
        origin = CellAddress.from_any(origin)
        for i, row in enumerate(rows, start=origin.row):
            for j, value in enumerate(row,):
                ws.cell(row=i, column=origin.col+j).value = value
        return ws
    except Exception as e:
        raise RuntimeError(LogMsg.err(e,
            rows=lute.annotate(rows),
            origin=str(origin)
        ))
    
