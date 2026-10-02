"""
*`@file`* `engines/pyxl/worksheet/read`
"""
from __future__ import annotations
from pathlib import Path
from typing import (Any, Optional, Mapping, Callable, Sequence, OrderedDict, overload)
import lute
from pyqi import (Array, String, Function)
from openpyxl.utils import (get_column_letter)
from noxl.models import (CellAddress, CellAddressLike, HeaderIndex, normalize_cell_value)
from ..types import (PyxlWorkbook, PyxlWorksheet, PyxlWorkbookSource, Row)
from ..utils import (resolve_pyxl_wb, get_visible_sheetnames)

LogMsg = lute.LogMsg
mlog = lute.get_logger(Path(__file__).parent.stem)

# kinda unnecessary but useful if don't want to type ws.cell(...).value all the time 
def cell_value(ws: PyxlWorksheet, /, row: int, column: int) -> Any:
    """ 
    *`@returns`* `ws.cell(row, column).value`
    """
    return ws.cell(row, column).value

def find_header(
    wb_source: PyxlWorkbookSource,
    sheetname: Optional[str | Callable[[str], bool]]=None,
    /,
    required_headers: Optional[Sequence[str]]=None,
    min_header_len: int = 2,
    scan_rows: int=25,
    scan_cols: int=50,
    unique_sheetname: bool=False,
    normalize_header_value: Optional[Callable[..., str]]=None,
    col_offset: Optional[int]=None
):
    wb: Optional[PyxlWorkbook] = None
    resolved_sheetname: Optional[str] = None
    try:
        if (scan_cols < 1 or scan_rows < 1):
            raise ValueError(
                "scan_rows and scan_cols must be positive integers"
            )
        if (not required_headers and min_header_len < 1):
            raise ValueError(
                "min_header_len must be positive integer when required_headers not provided"
            )
        if (col_offset and (not isinstance(col_offset, int) or col_offset < 0)):
            raise ValueError(LogMsg.bad_arg(
                col_offset, "col_offset", "Nonnegative integer required", int
            ))
        if (col_offset is None): col_offset = 0
        wb = resolve_pyxl_wb(wb_source)
        valid_sheetnames = wb.sheetnames
        if (isinstance(sheetname, str)):
            if (sheetname not in valid_sheetnames):
                raise KeyError(LogMsg(
                    f"sheetname not found in workbook",
                    workbook=wb.path,
                    sheetname=sheetname,
                    valid_sheetnames=valid_sheetnames
                ))
            resolved_sheetname = sheetname
        elif isinstance(sheetname, Callable):
            resolved_sheetname = Array.find(
                valid_sheetnames, 
                sheetname, 
                unique=unique_sheetname, 
                strict=True
            )
        else:
            resolved_sheetname = get_visible_sheetnames(wb)[0]
        ws = wb[resolved_sheetname]
        max_scan_row = min(scan_rows, ws.max_row)
        max_scan_col = min(scan_cols, ws.max_column)
        _rows = get_all_rows(ws)
        def acc_rowdict(acc: OrderedDict[str, int], curr: Any, i: int):
            if (not curr): return acc
            value = str(curr).strip()
            if (not value): return acc
            acc[value] = i
            return acc
        if (normalize_header_value is None):
            normalize_header_value = Function.Partial(
                normalize_cell_value, 
                case_sensitive=False
            )
        for row_idx, r in enumerate(_rows[:max_scan_row], start=1):
            if row_idx > scan_rows:
                break
            row = r[col_offset:max_scan_col]
            rowdict = Array.reduce(row, acc_rowdict, OrderedDict(), 1+col_offset)
            header_found: bool = False
            norm_row = set(Array.map(rowdict.keys(), normalize_header_value))
            if (required_headers):
                header_found = Array.every(
                    Array.map(required_headers, normalize_header_value),
                    lambda h: h in norm_row
                )
            else:
                header_found = len(norm_row) >= min_header_len
            
            if (header_found):
                return HeaderIndex(row_idx, columns=rowdict)
            
        raise RuntimeError(LogMsg(
            f"Unable to find header in target worksheet",
            max_scan_row=max_scan_row,
            max_scan_col=max_scan_col
        ))
    except Exception as e:
        raise RuntimeError(LogMsg.err(
            e,
            scan_rows=scan_rows,
            scan_cols=scan_cols,
            wb_source=lute.annotate(wb_source),
            resolved_sheetname=resolved_sheetname,
            required_headers=required_headers,
            min_header_len=min_header_len
        )) from e
    finally:
        if wb is not None: wb.close()


def get_all_rows(ws: PyxlWorksheet, /, values_only: bool=True) -> list[Row]:
    """ *`@returns`* `list(ws.iter_rows(values_only=values_only)) = list[Sequence[Any]]` """
    return list(ws.iter_rows(values_only=values_only))

@overload
def rows(
    src: PyxlWorkbookSource,
    /, 
    sheetname: str,
    origin: CellAddressLike="A1", 
    max_cols: Optional[int]=None, 
    values_only: bool=True
) -> list[Row]: ...

@overload
def rows(
    src: PyxlWorksheet,
    *, 
    values_only: bool=True,
) -> list[Row]: ...

def rows(
    src: PyxlWorksheet | PyxlWorkbookSource,
    sheetname: Optional[str]=None,  
    origin: Optional[CellAddressLike]=None, 
    max_cols: Optional[int]=None, 
    values_only: bool=True
) -> list[Row]:
    wb: Optional[PyxlWorkbook]=None
    try:
        ws: Optional[PyxlWorksheet]=None
        if (isinstance(src, PyxlWorksheet)):
            ws = src
        else:
            if (not sheetname): 
                raise ValueError(LogMsg.bad_arg(
                    sheetname, "sheetname", 
                    "sheetname must be defined when src is PyxlWorkbookSource", 
                    str
                ))
            wb = resolve_pyxl_wb(src)
            ws = wb[sheetname]
        origin = CellAddress.from_any(origin) if origin else CellAddress(1, 1)
        result: list[Row] = []
        row_start = max(origin.row, ws.min_row)
        col_start = max(origin.col, ws.min_column)
        col_end = (
            origin.col + max_cols - 1 if max_cols 
            else ws.max_column
        )
        for row in ws.iter_rows(
            min_row=row_start,
            min_col=col_start,
            max_col=col_end,
            values_only=values_only
        ):
            result.append(row)
        return result
    except Exception as e:
        raise RuntimeError(LogMsg.err(
            e,
            src=lute.annotate(src),
            sheetname=sheetname,
            wb="<unknown>" if wb is None else wb.path,
            origin=str(origin),
            max_cols=max_cols,
            values_only=values_only
        ))
    finally:
        if (wb is not None): wb.close()