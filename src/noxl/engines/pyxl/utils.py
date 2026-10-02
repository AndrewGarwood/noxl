"""
*`@file`* `engines/pyxl/utils`
"""

from typing import Any
from pyqi import Array
from openpyxl import load_workbook
from .types import (PyxlWorkbook, PyxlWorkbookSource, PyxlWorksheet, PyxlSheetInfo)

def resolve_pyxl_wb(
    wb_source: PyxlWorkbookSource,
    /, 
    read_only: bool = False, 
    data_only: bool = True
) -> PyxlWorkbook:
    return (
        wb_source if isinstance(wb_source, PyxlWorkbook)
        else load_workbook(wb_source, read_only=read_only, data_only=data_only)
    )

def is_hidden_worksheet(ws: PyxlWorksheet, /):
    return "hidden" in ws.sheet_state.lower()

def get_visible_sheetnames(wb_source: PyxlWorkbookSource, /) -> list[str]:
    wb = resolve_pyxl_wb(wb_source)
    return Array.filter(wb.sheetnames, lambda k: not is_hidden_worksheet(wb[k]))

def get_sheetinfo(wb_source: PyxlWorkbookSource, /) -> dict[str, PyxlSheetInfo]:
    wb = resolve_pyxl_wb(wb_source, True, True)
    result: dict[str, PyxlSheetInfo] = dict()
    for i, s in enumerate(wb.sheetnames):
        result[s] = PyxlSheetInfo(wb[s].sheet_state, i)
    if not isinstance(wb_source, PyxlWorkbook):
        wb.close()
    return result
